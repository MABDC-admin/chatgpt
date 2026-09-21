"""Render generated Markdown into downloadable DOCX, PDF and PPTX files.

The teacher tools stream Markdown. Teachers need something they can print,
hand out, or drop into a staff folder, so this turns that Markdown into a
real document rather than asking them to copy-paste into Word.

Everything is produced in memory and returned as bytes; nothing is written
to disk, so there is no cleanup or storage growth to manage.
"""

import io
import re

# Markdown we actually emit from the tool prompts: headings, bullets,
# numbered items, bold and italic. Anything else is treated as body text.
_H2 = re.compile(r"^##\s+(.*)$")
_H3 = re.compile(r"^###\s+(.*)$")
_BULLET = re.compile(r"^[-*]\s+(.*)$")
_NUMBERED = re.compile(r"^(\d+)[.)]\s+(.*)$")
_BOLD = re.compile(r"\*\*(.+?)\*\*")
_ITALIC = re.compile(r"(?<!\*)\*([^*]+?)\*(?!\*)")
_RULE = re.compile(r"^-{3,}$")


def _strip_marks(text: str) -> str:
    """Remove inline markdown so plain-text renderers do not show the syntax."""
    return _ITALIC.sub(r"\1", _BOLD.sub(r"\1", text))


class Block:
    __slots__ = ("kind", "text")

    def __init__(self, kind: str, text: str):
        self.kind = kind
        self.text = text


def parse(markdown: str) -> list[Block]:
    """Flatten Markdown into a list of typed blocks the renderers can walk."""
    blocks: list[Block] = []
    for raw in markdown.splitlines():
        line = raw.rstrip()
        if not line.strip():
            continue
        if _RULE.match(line.strip()):
            blocks.append(Block("rule", ""))
            continue
        if m := _H3.match(line):
            blocks.append(Block("h3", _strip_marks(m.group(1))))
        elif m := _H2.match(line):
            blocks.append(Block("h2", _strip_marks(m.group(1))))
        elif m := _BULLET.match(line):
            blocks.append(Block("bullet", _strip_marks(m.group(1))))
        elif m := _NUMBERED.match(line):
            blocks.append(Block("numbered", f"{m.group(1)}. {_strip_marks(m.group(2))}"))
        else:
            blocks.append(Block("body", _strip_marks(line)))
    return blocks


def to_docx(markdown: str, title: str) -> bytes:
    import docx
    from docx.shared import Pt

    document = docx.Document()
    document.add_heading(title, level=0)

    for block in parse(markdown):
        if block.kind == "h2":
            document.add_heading(block.text, level=1)
        elif block.kind == "h3":
            document.add_heading(block.text, level=2)
        elif block.kind == "bullet":
            document.add_paragraph(block.text, style="List Bullet")
        elif block.kind == "numbered":
            document.add_paragraph(block.text, style="List Number")
        elif block.kind == "rule":
            document.add_paragraph("_" * 50)
        else:
            para = document.add_paragraph(block.text)
            para.paragraph_format.space_after = Pt(6)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def to_pdf(markdown: str, title: str) -> bytes:
    from reportlab.lib.enums import TA_LEFT
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        HRFlowable,
        ListFlowable,
        ListItem,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
    )
    from xml.sax.saxutils import escape

    buffer = io.BytesIO()
    document = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        title=title,
        leftMargin=20 * mm,
        rightMargin=20 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
    )

    sheet = getSampleStyleSheet()
    body = ParagraphStyle(
        "Body", parent=sheet["Normal"], fontSize=10.5, leading=15,
        alignment=TA_LEFT, spaceAfter=5,
    )
    h1 = ParagraphStyle("H1", parent=sheet["Heading1"], fontSize=17, spaceAfter=10)
    h2 = ParagraphStyle("H2", parent=sheet["Heading2"], fontSize=13, spaceBefore=11, spaceAfter=6)
    h3 = ParagraphStyle("H3", parent=sheet["Heading3"], fontSize=11.5, spaceBefore=8, spaceAfter=4)

    story: list = [Paragraph(escape(title), h1), Spacer(1, 4)]
    pending: list = []

    def flush_list() -> None:
        """Group consecutive bullets into one ListFlowable for correct indents."""
        if pending:
            story.append(ListFlowable(pending[:], bulletType="bullet", leftIndent=14))
            story.append(Spacer(1, 4))
            pending.clear()

    for block in parse(markdown):
        if block.kind == "bullet":
            pending.append(ListItem(Paragraph(escape(block.text), body)))
            continue
        flush_list()
        if block.kind == "h2":
            story.append(Paragraph(escape(block.text), h2))
        elif block.kind == "h3":
            story.append(Paragraph(escape(block.text), h3))
        elif block.kind == "rule":
            story.append(Spacer(1, 6))
            story.append(HRFlowable(width="100%", thickness=0.6, color="#999999"))
            story.append(Spacer(1, 6))
        else:
            story.append(Paragraph(escape(block.text), body))

    flush_list()
    document.build(story)
    return buffer.getvalue()


def to_pptx(slides: list[dict], title: str, subtitle: str = "") -> bytes:
    """Build a deck from [{title, bullets: [...], notes: str}, ...]."""
    from pptx import Presentation
    from pptx.util import Pt

    deck = Presentation()

    cover = deck.slides.add_slide(deck.slide_layouts[0])
    cover.shapes.title.text = title
    if len(cover.placeholders) > 1:
        cover.placeholders[1].text = subtitle

    for slide_spec in slides:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = str(slide_spec.get("title", ""))[:120]

        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = [str(b) for b in slide_spec.get("bullets", []) if str(b).strip()]
        for i, bullet in enumerate(bullets):
            # text_frame always starts with one empty paragraph; fill it before
            # adding more, otherwise every deck opens with a blank first line.
            para = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            para.text = bullet
            para.level = 0
            para.font.size = Pt(18)

        notes = str(slide_spec.get("notes", "")).strip()
        if notes:
            slide.notes_slide.notes_text_frame.text = notes

    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def safe_filename(name: str, extension: str) -> str:
    """Make a filename that survives Content-Disposition and every OS."""
    cleaned = re.sub(r"[^A-Za-z0-9 _-]", "", name).strip() or "document"
    return f"{cleaned[:60].replace(' ', '_')}.{extension}"
