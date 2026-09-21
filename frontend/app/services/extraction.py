"""Turn an uploaded file into text the model can read.

Images are not extracted here -- they are passed to the model as vision input
by the chat router, which needs the bytes rather than a transcript.
"""

import csv
import io
from pathlib import Path

MAX_CHARS_PER_FILE = 40_000

IMAGE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif"}

TEXT_TYPES = {
    "text/plain": "txt",
    "text/markdown": "txt",
    "text/csv": "csv",
    "application/json": "txt",
}

DOC_TYPES = {
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet": "xlsx",
}

SUFFIX_FALLBACK = {
    ".pdf": "pdf", ".docx": "docx", ".pptx": "pptx", ".xlsx": "xlsx",
    ".txt": "txt", ".md": "txt", ".csv": "csv", ".json": "txt",
    ".png": "image", ".jpg": "image", ".jpeg": "image", ".webp": "image", ".gif": "image",
}

ACCEPTED = set(IMAGE_TYPES) | set(TEXT_TYPES) | set(DOC_TYPES)


def kind_of(mime: str, filename: str) -> str:
    """Classify an upload. Browsers send inconsistent MIME types, so the file
    suffix is the tie-breaker rather than a hard failure."""
    if mime in IMAGE_TYPES:
        return "image"
    if mime in DOC_TYPES:
        return DOC_TYPES[mime]
    if mime in TEXT_TYPES:
        return TEXT_TYPES[mime]
    return SUFFIX_FALLBACK.get(Path(filename).suffix.lower(), "unsupported")


def _truncate(text: str) -> str:
    text = text.strip()
    if len(text) <= MAX_CHARS_PER_FILE:
        return text
    return text[:MAX_CHARS_PER_FILE] + "\n\n[... truncated, file continues ...]"


def _pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        body = (page.extract_text() or "").strip()
        if body:
            pages.append(f"--- Page {i} ---\n{body}")
    return "\n\n".join(pages)


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts = [p.text for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def _pptx(data: bytes) -> str:
    from pptx import Presentation

    deck = Presentation(io.BytesIO(data))
    slides = []
    for i, slide in enumerate(deck.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                lines.append(shape.text_frame.text.strip())
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
            note_text = slide.notes_slide.notes_text_frame.text.strip()
            if note_text:
                notes = f"\n[Speaker notes] {note_text}"
        if lines or notes:
            slides.append(f"--- Slide {i} ---\n" + "\n".join(lines) + notes)
    return "\n\n".join(slides)


def _xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheets = []
    for sheet in book.worksheets:
        rows = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            sheets.append(f"--- Sheet: {sheet.title} ---\n" + "\n".join(rows))
    book.close()
    return "\n\n".join(sheets)


def _csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    rows = list(csv.reader(io.StringIO(text)))
    return "\n".join(" | ".join(r) for r in rows if any(r))


def extract(data: bytes, kind: str) -> str:
    """Extract text. Raises ValueError with a readable reason on failure."""
    try:
        if kind == "pdf":
            text = _pdf(data)
        elif kind == "docx":
            text = _docx(data)
        elif kind == "pptx":
            text = _pptx(data)
        elif kind == "xlsx":
            text = _xlsx(data)
        elif kind == "csv":
            text = _csv(data)
        elif kind == "txt":
            text = data.decode("utf-8", errors="replace")
        else:
            return ""
    except Exception as exc:
        raise ValueError(f"Could not read this {kind.upper()} file: {exc}") from exc

    if not text.strip() and kind == "pdf":
        raise ValueError(
            "No text found in this PDF. It is probably a scan; OCR is not enabled yet."
        )
    return _truncate(text)
